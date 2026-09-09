import sqlite3
import os
import uuid
import csv
import io
import hmac
from werkzeug.utils import secure_filename
from flask import Flask, render_template, request, redirect, url_for, flash, session, jsonify, send_file, make_response
from werkzeug.security import generate_password_hash, check_password_hash
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.pdfgen import canvas
from init_db import create_tables, migrate_files_to_db
from push import save_subscription, send_push_to_all, VAPID_PUBLIC_KEY, push_configured

app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY', 'your_secret_key')
app.config['MAX_CONTENT_LENGTH'] = 8 * 1024 * 1024  # 8 MB upload limit

DB_FILE = os.environ.get('DB_PATH', 'votestack3.db')

# ── uploads ────────────────────────────────────────────────────────────────
UPLOAD_ROOT      = os.path.join('static', 'uploads')
ALLOWED_UPLOAD_EXTS = {'png', 'jpg', 'jpeg', 'gif', 'webp'}

for sub in ('announcements', 'events', 'gallery'):
    os.makedirs(os.path.join(UPLOAD_ROOT, sub), exist_ok=True)


def allowed_image(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_UPLOAD_EXTS


def save_uploaded_image(file_storage, subfolder):
    """Saves an uploaded image under static/uploads/<subfolder>/ and
    returns the web path to store in the DB, or None if no valid file."""
    if not file_storage or file_storage.filename == '':
        return None
    if not allowed_image(file_storage.filename):
        flash("Only image files (png, jpg, jpeg, gif, webp) are allowed.", "danger")
        return None

    ext      = file_storage.filename.rsplit('.', 1)[1].lower()
    filename = f"{uuid.uuid4().hex}.{ext}"
    filename = secure_filename(filename)
    rel_path = os.path.join(UPLOAD_ROOT, subfolder, filename)
    file_storage.save(rel_path)
    return '/' + rel_path.replace(os.sep, '/')

# ── file paths kept for backward-compat / PDF generation ──────────────────────
usernames_file = 'usernames.txt'
choices_file   = 'choices.txt'
votes_file     = 'votes.txt'
role_file      = 'roles.txt'

# ── initialise DB on every startup ────────────────────────────────────────────
with app.app_context():
    create_tables(DB_FILE)
    migrate_files_to_db(DB_FILE, role_file, choices_file, votes_file)

# ── helpers ───────────────────────────────────────────────────────────────────
def get_db():
    conn = sqlite3.connect(DB_FILE)
    conn.row_factory = sqlite3.Row
    return conn

def load_role():
    """Return the current role from the DB."""
    conn = get_db()
    row = conn.execute(
        "SELECT value FROM settings WHERE key='current_role'"
    ).fetchone()
    conn.close()
    return row['value'] if row else None

def load_voting_status():
    """Return 'open' or 'closed' (defaults to 'closed' if not set)."""
    conn = get_db()
    row = conn.execute(
        "SELECT value FROM settings WHERE key='voting_status'"
    ).fetchone()
    conn.close()
    return row['value'] if row else 'closed'

# ── auth ──────────────────────────────────────────────────────────────────────
@app.route('/admin_login', methods=['GET', 'POST'])
def admin_login():
    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '')

        if not username or not password:
            flash("Please enter both username and password.", "danger")
            return redirect(url_for('admin_login'))

        try:
            conn = get_db()
            admin = conn.execute(
                "SELECT * FROM admins WHERE username=?", (username,)
            ).fetchone()
            conn.close()

            if admin and check_password_hash(admin['password_hash'], password):
                session.clear()
                session['username'] = 'admin'
                session['is_admin'] = True
                return redirect(url_for('admin_dashboard'))
            else:
                flash("Invalid username or password.", "danger")
                return redirect(url_for('admin_login'))

        except Exception as e:
            flash(f"Database error: {e}", "danger")
            return redirect(url_for('admin_login'))

    return render_template('admin_login.html')


@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'GET':
        if 'username' in session:
            return redirect(url_for('vote'))
        next_url = request.args.get('next', '') or url_for('vote')
        return render_template('login.html', next=next_url)

    session.clear()
    phone    = request.form.get('username', '').strip()
    phone    = ''.join(phone.split())
    next_url = request.form.get('next', '').strip()

    # only allow relative redirects, never an external URL
    if not next_url.startswith('/'):
        next_url = url_for('vote')

    if not phone:
        flash("Please enter a phone number.", "danger")
        return redirect(url_for('login', next=next_url))

    try:
        conn = get_db()
        user = conn.execute(
            "SELECT * FROM users WHERE phone_number=?", (phone,)
        ).fetchone()
        conn.close()

        if user:
            session['username'] = phone
            flash(f"Welcome!", "success")
            return redirect(next_url)
        else:
            flash("Phone number not authorized.", "danger")
            return redirect(url_for('login', next=next_url))

    except Exception as e:
        flash(f"Database error: {e}", "danger")
        return redirect(url_for('login', next=next_url))


@app.route('/logout')
def logout():
    session.clear()
    flash("You have been logged out.", "success")
    return redirect(url_for('index'))


@app.route('/setup_admin', methods=['GET', 'POST'])
def setup_admin():
    conn = get_db()
    existing = conn.execute("SELECT id FROM admins LIMIT 1").fetchone()
    conn.close()

    if existing:
        flash("Admin account already exists.", "warning")
        return redirect(url_for('admin_login'))

    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '')
        confirm  = request.form.get('confirm', '')

        if not username or not password:
            flash("All fields are required.", "danger")
            return redirect(url_for('setup_admin'))
        if password != confirm:
            flash("Passwords do not match.", "danger")
            return redirect(url_for('setup_admin'))
        if len(password) < 8:
            flash("Password must be at least 8 characters.", "danger")
            return redirect(url_for('setup_admin'))

        conn = get_db()
        conn.execute(
            "INSERT INTO admins (username, password_hash) VALUES (?, ?)",
            (username, generate_password_hash(password))
        )
        conn.commit()
        conn.close()
        flash("Admin account created! Please log in.", "success")
        return redirect(url_for('admin_login'))

    return render_template('setup_admin.html')

# ── public pages ──────────────────────────────────────────────────────────────
@app.route('/', methods=['GET'])
def index():
    return redirect(url_for('announcements'))


@app.route('/rsvp', methods=['GET', 'POST'])
@app.route('/admin/rsvp')
def admin_rsvp():
    if not admin_required():
        flash('Access restricted to admin only.', 'danger')
        return redirect(url_for('admin_login'))
    conn = get_db()
    rows = conn.execute('SELECT * FROM rsvps ORDER BY attending DESC, first_name COLLATE NOCASE').fetchall()
    conn.close()
    attending_count = sum(1 for r in rows if r['attending'])
    return render_template('admin_rsvp.html', rsvps=rows, attending_count=attending_count, not_attending_count=len(rows)-attending_count, total_count=len(rows))

@app.route('/admin/rsvp/download')
def download_rsvps():
    if not admin_required():
        flash('Access restricted to admin only.', 'danger')
        return redirect(url_for('admin_login'))
    conn = get_db()
    rows = conn.execute('SELECT first_name, attending, created_at, updated_at FROM rsvps ORDER BY first_name COLLATE NOCASE').fetchall()
    conn.close()
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(['First Name', 'Attending', 'Created At', 'Updated At'])
    for row in rows:
        writer.writerow([row['first_name'], 'Yes' if row['attending'] else 'No', row['created_at'], row['updated_at']])
    response = make_response(output.getvalue())
    response.headers['Content-Type'] = 'text/csv; charset=utf-8'
    response.headers['Content-Disposition'] = 'attachment; filename=practical_evangelism_rsvps.csv'
    return response


@app.route('/vote')
def vote():
    if 'username' not in session:
        return redirect(url_for('login', next='/vote'))

    role   = load_role()
    status = load_voting_status()

    conn = get_db()
    rows = conn.execute(
        "SELECT choice FROM choices WHERE role=?", (role,)
    ).fetchall() if role else []

    has_voted = False
    tally     = {}
    if role:
        existing = conn.execute(
            "SELECT id FROM votes WHERE username=? AND role=?",
            (session['username'], role)
        ).fetchone()
        has_voted = existing is not None

        if status == 'closed':
            tally_rows = conn.execute(
                "SELECT choice, COUNT(*) as cnt FROM votes WHERE role=? GROUP BY choice",
                (role,)
            ).fetchall()
            tally = {r['choice']: r['cnt'] for r in tally_rows}

    conn.close()
    choices = [r['choice'] for r in rows]

    return render_template(
        'vote.html',
        role=role,
        status=status,
        choices=choices,
        has_voted=has_voted,
        tally=tally,
    )


@app.route('/view_role')
def view_role():
    role = load_role()
    return render_template('view_role.html', role=role)

# ── announcements (public) ──────────────────────────────────────────────────
@app.route('/announcements')
def announcements():
    conn = get_db()
    rows = conn.execute(
        "SELECT * FROM announcements ORDER BY created_at DESC"
    ).fetchall()
    conn.close()
    return render_template('announcements.html', announcements=rows)

# ── events (public) ─────────────────────────────────────────────────────────
@app.route('/events')
def events():
    conn = get_db()
    upcoming = conn.execute(
        "SELECT * FROM events WHERE event_date >= date('now') ORDER BY event_date ASC, event_time ASC"
    ).fetchall()
    past = conn.execute(
        "SELECT * FROM events WHERE event_date < date('now') ORDER BY event_date DESC, event_time DESC"
    ).fetchall()
    conn.close()
    return render_template('events.html', upcoming=upcoming, past=past)

# ── gallery (public) ────────────────────────────────────────────────────────
@app.route('/gallery')
def gallery():
    conn = get_db()
    rows = conn.execute(
        "SELECT * FROM gallery ORDER BY created_at DESC"
    ).fetchall()
    conn.close()
    return render_template('gallery.html', photos=rows)

# ── push notifications ──────────────────────────────────────────────────────
@app.route('/vapid_public_key')
def vapid_public_key():
    return jsonify({"publicKey": VAPID_PUBLIC_KEY, "configured": push_configured()})


@app.route('/save_subscription', methods=['POST'])
def save_subscription_route():
    # No login required — announcements/events are public, so anyone who
    # installs the app can opt in to notifications, not just voters.
    subscription = request.get_json(silent=True)
    if not subscription:
        return jsonify({"error": "Missing subscription payload"}), 400
    ok = save_subscription(DB_FILE, subscription, username=session.get('username'))
    if not ok:
        return jsonify({"error": "Invalid subscription payload"}), 400
    return jsonify({"message": "Subscribed"}), 200

# ── voting ────────────────────────────────────────────────────────────────────
@app.route('/submit_vote', methods=['POST'])
def submit_vote():
    if 'username' not in session:
        return redirect(url_for('login', next='/vote'))

    role   = load_role()
    status = load_voting_status()

    if not role:
        flash("No active role to vote on.", "danger")
        return redirect(url_for('vote'))

    if status != 'open':
        flash("Voting is currently closed.", "danger")
        return redirect(url_for('vote'))

    choice = request.form.get('choice')
    if not choice:
        flash("Please select a choice to vote for.", "danger")
        return redirect(url_for('vote'))

    username = session['username']
    try:
        conn = get_db()
        existing = conn.execute(
            "SELECT id FROM votes WHERE username=? AND role=?", (username, role)
        ).fetchone()

        if existing:
            flash(f"You have already voted for '{role}'.", "danger")
        else:
            conn.execute(
                "INSERT INTO votes (username, role, choice) VALUES (?, ?, ?)",
                (username, role, choice)
            )
            conn.commit()
            flash(f"Your vote for {choice} has been recorded!", "success")
        conn.close()

    except Exception as e:
        flash(f"An error occurred: {e}", "danger")

    return redirect(url_for('vote'))

# ── API endpoints ─────────────────────────────────────────────────────────────
@app.route('/api/current_role', methods=['GET'])
def get_current_role_api():
    return jsonify({"role": load_role()}), 200


@app.route('/api/voter_status', methods=['GET'])
def voter_status():
    if 'username' not in session:
        return jsonify({"error": "Not logged in"}), 401

    role   = load_role()
    status = load_voting_status()

    if not role:
        return jsonify({"role": None, "status": "closed", "choices": [], "has_voted": False, "tally": {}})

    conn = get_db()
    rows = conn.execute(
        "SELECT choice FROM choices WHERE role=?", (role,)
    ).fetchall()
    existing = conn.execute(
        "SELECT id FROM votes WHERE username=? AND role=?",
        (session['username'], role)
    ).fetchone()

    tally = {}
    if status == 'closed':
        tally_rows = conn.execute(
            "SELECT choice, COUNT(*) as cnt FROM votes WHERE role=? GROUP BY choice",
            (role,)
        ).fetchall()
        tally = {r['choice']: r['cnt'] for r in tally_rows}

    conn.close()

    return jsonify({
        "role":      role,
        "status":    status,
        "choices":   [r['choice'] for r in rows],
        "has_voted": existing is not None,
        "tally":     tally,
    })


@app.route('/api/submit_vote', methods=['POST'])
def submit_vote_api():
    if 'username' not in session:
        return jsonify({"error": "Please log in first."}), 401

    role   = load_role()
    status = load_voting_status()

    if status != 'open':
        return jsonify({"error": "Voting is currently closed."}), 400

    data   = request.get_json()
    choice = data.get('choice') if data else None
    if not choice:
        return jsonify({"error": "Please select a choice."}), 400

    username = session['username']
    try:
        conn = get_db()
        existing = conn.execute(
            "SELECT id FROM votes WHERE username=? AND role=?", (username, role)
        ).fetchone()

        if existing:
            conn.close()
            return jsonify({"error": "You have already voted."}), 400

        conn.execute(
            "INSERT INTO votes (username, role, choice) VALUES (?, ?, ?)",
            (username, role, choice)
        )
        conn.commit()
        conn.close()
        return jsonify({"message": f"Your vote for {choice} has been recorded!"}), 200

    except Exception as e:
        return jsonify({"error": str(e)}), 500

# ── admin dashboard ───────────────────────────────────────────────────────────
def admin_required():
    return session.get('is_admin') is True


@app.route('/admin_dashboard', methods=['GET'])
def admin_dashboard():
    if not admin_required():
        flash("Access restricted to admin only.", "danger")
        return redirect(url_for('index'))
    role        = load_role()
    voting_open = load_voting_status() == 'open'
    return render_template('admin_dashboard.html', role=role, voting_open=voting_open)


@app.route('/toggle_voting', methods=['POST'])
def toggle_voting():
    if not admin_required():
        flash("Access restricted to admin only.", "danger")
        return redirect(url_for('index'))

    current    = load_voting_status()
    new_status = 'closed' if current == 'open' else 'open'

    conn = get_db()
    conn.execute(
        "INSERT OR REPLACE INTO settings (key, value) VALUES ('voting_status', ?)",
        (new_status,)
    )
    conn.commit()
    conn.close()

    if new_status == 'open':
        flash("Voting is now open. Voters can cast their ballots.", "success")
    else:
        flash("Voting is now closed. Tally is visible to all voters.", "success")

    return redirect(url_for('admin_dashboard'))


@app.route('/update_role', methods=['POST'])
def update_role():
    if not admin_required():
        flash("Access restricted to admin only.", "danger")
        return redirect(url_for('index'))

    if load_voting_status() == 'open':
        flash("Close voting before changing the role.", "danger")
        return redirect(url_for('admin_dashboard'))

    new_role = request.form.get('new_role', '').strip()
    if not new_role:
        flash("Role cannot be empty.", "danger")
        return redirect(url_for('admin_dashboard'))

    conn = get_db()
    conn.execute(
        "INSERT OR REPLACE INTO settings (key, value) VALUES ('current_role', ?)",
        (new_role,)
    )
    conn.commit()
    conn.close()

    with open(role_file, 'w') as f:
        f.write(new_role)

    flash(f"Role updated to '{new_role}'!", "success")
    return redirect(url_for('admin_dashboard'))


@app.route('/update_choices', methods=['POST'])
def update_choices():
    if not admin_required():
        flash("Access restricted to admin only.", "danger")
        return redirect(url_for('index'))

    if load_voting_status() == 'open':
        flash("Close voting before changing choices.", "danger")
        return redirect(url_for('admin_dashboard'))

    role = load_role()
    if not role:
        flash("Set a role first before adding choices.", "warning")
        return redirect(url_for('admin_dashboard'))

    raw         = request.form.get('choices', '')
    new_choices = [c.strip() for c in raw.splitlines() if c.strip()]

    try:
        conn = get_db()
        conn.execute("DELETE FROM choices WHERE role=?", (role,))
        conn.executemany(
            "INSERT OR IGNORE INTO choices (role, choice) VALUES (?, ?)",
            [(role, c) for c in new_choices]
        )
        conn.commit()
        conn.close()
        flash("Choices updated successfully!", "success")
    except Exception as e:
        flash(f"Error updating choices: {e}", "danger")

    return redirect(url_for('admin_dashboard'))


@app.route('/view_choices', methods=['GET'])
def view_choices():
    if not admin_required():
        flash("Access restricted to admin only.", "danger")
        return redirect(url_for('index'))

    role = load_role()
    conn = get_db()
    rows = conn.execute("SELECT choice FROM choices WHERE role=?", (role,)).fetchall() if role else []
    conn.close()
    choices = [r['choice'] for r in rows]
    return render_template('view_choices.html', choices=choices)


@app.route('/generate_tally', methods=['GET', 'POST'])
def generate_tally():
    if not admin_required():
        flash("Access restricted to admin only.", "danger")
        return redirect(url_for('index'))

    role = load_role()
    if not role:
        flash("No role is currently set.", "warning")
        return redirect(url_for('admin_dashboard'))

    conn = get_db()
    rows = conn.execute(
        "SELECT choice, COUNT(*) as cnt FROM votes WHERE role=? GROUP BY choice",
        (role,)
    ).fetchall()
    conn.close()
    tally = {r['choice']: r['cnt'] for r in rows}
    return render_template('tally.html', role=role, tally=tally)

# ── admin: announcements ────────────────────────────────────────────────────
@app.route('/admin/announcements')
def admin_announcements():
    if not admin_required():
        flash("Access restricted to admin only.", "danger")
        return redirect(url_for('index'))
    conn = get_db()
    rows = conn.execute("SELECT * FROM announcements ORDER BY created_at DESC").fetchall()
    conn.close()
    return render_template('admin_announcements.html', announcements=rows)


@app.route('/admin/announcements/create', methods=['POST'])
def create_announcement():
    if not admin_required():
        flash("Access restricted to admin only.", "danger")
        return redirect(url_for('index'))

    title = request.form.get('title', '').strip()
    body  = request.form.get('body', '').strip()
    if not title or not body:
        flash("Title and message are required.", "danger")
        return redirect(url_for('admin_announcements'))

    image_path = save_uploaded_image(request.files.get('image'), 'announcements')

    conn = get_db()
    conn.execute(
        "INSERT INTO announcements (title, body, image_path) VALUES (?, ?, ?)",
        (title, body, image_path)
    )
    conn.commit()
    conn.close()

    result = send_push_to_all(DB_FILE, title=f"📢 {title}", body=body[:120], url='/announcements')
    if result.get('skipped'):
        flash("Announcement posted! (Push notifications aren't configured yet.)", "success")
    else:
        flash(f"Announcement posted and pushed to {result['sent']} device(s).", "success")

    return redirect(url_for('admin_announcements'))


@app.route('/admin/announcements/delete/<int:announcement_id>', methods=['POST'])
def delete_announcement(announcement_id):
    if not admin_required():
        flash("Access restricted to admin only.", "danger")
        return redirect(url_for('index'))
    conn = get_db()
    conn.execute("DELETE FROM announcements WHERE id=?", (announcement_id,))
    conn.commit()
    conn.close()
    flash("Announcement deleted.", "success")
    return redirect(url_for('admin_announcements'))

# ── admin: events ────────────────────────────────────────────────────────────
@app.route('/admin/events')
def admin_events():
    if not admin_required():
        flash("Access restricted to admin only.", "danger")
        return redirect(url_for('index'))
    conn = get_db()
    rows = conn.execute("SELECT * FROM events ORDER BY event_date DESC").fetchall()
    conn.close()
    return render_template('admin_events.html', events=rows)


@app.route('/admin/events/create', methods=['POST'])
def create_event():
    if not admin_required():
        flash("Access restricted to admin only.", "danger")
        return redirect(url_for('index'))

    title       = request.form.get('title', '').strip()
    description = request.form.get('description', '').strip()
    location    = request.form.get('location', '').strip()
    event_date  = request.form.get('event_date', '').strip()
    event_time  = request.form.get('event_time', '').strip()

    if not title or not event_date:
        flash("Title and date are required.", "danger")
        return redirect(url_for('admin_events'))

    image_path = save_uploaded_image(request.files.get('image'), 'events')

    conn = get_db()
    conn.execute(
        """INSERT INTO events (title, description, location, event_date, event_time, image_path)
           VALUES (?, ?, ?, ?, ?, ?)""",
        (title, description, location, event_date, event_time, image_path)
    )
    conn.commit()
    conn.close()

    when = f"{event_date}" + (f" at {event_time}" if event_time else "")
    push_body = f"{when}" + (f" — {location}" if location else "")
    result = send_push_to_all(DB_FILE, title=f"📅 {title}", body=push_body[:120], url='/events')
    if result.get('skipped'):
        flash("Event posted! (Push notifications aren't configured yet.)", "success")
    else:
        flash(f"Event posted and pushed to {result['sent']} device(s).", "success")

    return redirect(url_for('admin_events'))


@app.route('/admin/events/delete/<int:event_id>', methods=['POST'])
def delete_event(event_id):
    if not admin_required():
        flash("Access restricted to admin only.", "danger")
        return redirect(url_for('index'))
    conn = get_db()
    conn.execute("DELETE FROM events WHERE id=?", (event_id,))
    conn.commit()
    conn.close()
    flash("Event deleted.", "success")
    return redirect(url_for('admin_events'))

# ── admin: gallery ───────────────────────────────────────────────────────────
@app.route('/admin/gallery')
def admin_gallery():
    if not admin_required():
        flash("Access restricted to admin only.", "danger")
        return redirect(url_for('index'))
    conn = get_db()
    rows = conn.execute("SELECT * FROM gallery ORDER BY created_at DESC").fetchall()
    conn.close()
    return render_template('admin_gallery.html', photos=rows)


@app.route('/admin/gallery/upload', methods=['POST'])
def upload_gallery_photo():
    if not admin_required():
        flash("Access restricted to admin only.", "danger")
        return redirect(url_for('index'))

    caption = request.form.get('caption', '').strip()
    files   = request.files.getlist('images')

    if not files or all(f.filename == '' for f in files):
        flash("Please choose at least one photo.", "danger")
        return redirect(url_for('admin_gallery'))

    conn   = get_db()
    added  = 0
    for f in files:
        image_path = save_uploaded_image(f, 'gallery')
        if image_path:
            conn.execute(
                "INSERT INTO gallery (caption, image_path) VALUES (?, ?)",
                (caption, image_path)
            )
            added += 1
    conn.commit()
    conn.close()

    if added:
        flash(f"{added} photo(s) added to the gallery.", "success")
    return redirect(url_for('admin_gallery'))


@app.route('/admin/gallery/delete/<int:photo_id>', methods=['POST'])
def delete_gallery_photo(photo_id):
    if not admin_required():
        flash("Access restricted to admin only.", "danger")
        return redirect(url_for('index'))
    conn = get_db()
    conn.execute("DELETE FROM gallery WHERE id=?", (photo_id,))
    conn.commit()
    conn.close()
    flash("Photo removed.", "success")
    return redirect(url_for('admin_gallery'))

# ── voter management ──────────────────────────────────────────────────────────
@app.route('/enter_voters', methods=['POST'])
def enter_voters():
    if not admin_required():
        flash("Access restricted to admin only.", "danger")
        return redirect(url_for('index'))

    raw    = request.form.get('phone_numbers', '')
    phones = [p.strip() for p in raw.splitlines() if p.strip()]

    if not phones:
        flash("Please enter at least one phone number.", "danger")
        return redirect(url_for('admin_dashboard'))

    try:
        conn    = get_db()
        added   = 0
        skipped = 0
        for p in phones:
            try:
                conn.execute("INSERT INTO users (phone_number) VALUES (?)", (p,))
                added += 1
            except sqlite3.IntegrityError:
                skipped += 1
        conn.commit()
        conn.close()

        msg = f"{added} voter(s) added."
        if skipped:
            msg += f" {skipped} duplicate(s) skipped."
        flash(msg, "success")

    except Exception as e:
        flash(f"Error: {e}", "danger")

    return redirect(url_for('admin_dashboard'))


@app.route('/view_usernames')
def view_usernames():
    if not admin_required():
        flash("Access restricted to admin only.", "danger")
        return redirect(url_for('index'))

    conn      = get_db()
    rows      = conn.execute("SELECT phone_number FROM users").fetchall()
    conn.close()
    usernames = [r['phone_number'] for r in rows]
    return render_template('view_usernames.html', usernames=usernames)


@app.route('/delete_all_voters', methods=['POST'])
def delete_all_voters():
    if not admin_required():
        flash("Access restricted to admin only.", "danger")
        return redirect(url_for('index'))

    conn = get_db()
    conn.execute("DELETE FROM users")
    conn.commit()
    conn.close()
    flash("All voters deleted.", "success")
    return redirect(url_for('view_usernames'))


@app.route('/delete_voter/<username>', methods=['POST'])
def delete_voter(username):
    if not admin_required():
        flash("Access restricted to admin only.", "danger")
        return redirect(url_for('index'))

    conn   = get_db()
    result = conn.execute("DELETE FROM users WHERE phone_number=?", (username,))
    conn.commit()
    conn.close()

    if result.rowcount:
        flash(f"Voter {username} removed.", "success")
    else:
        flash(f"Voter {username} not found.", "warning")

    return redirect(url_for('view_usernames'))


@app.route('/download_voters_pdf')
def download_voters_pdf():
    file_path = os.path.join('static', 'voters_list.pdf')
    if not os.path.exists(file_path):
        flash("Voters list PDF not found.", "danger")
        return redirect(url_for('admin_dashboard'))
    return send_file(file_path, as_attachment=True)

# ── document helpers ──────────────────────────────────────────────────────────
def create_word_document(usernames, output_file='static/voters_list.docx'):
    doc = Document()
    doc.add_heading('Voters List', level=1)
    for u in usernames:
        doc.add_paragraph(u)
    doc.save(output_file)


def generate_pdf(usernames, output_file='static/voters_list.pdf'):
    os.makedirs('static', exist_ok=True)
    c            = canvas.Canvas(output_file, pagesize=letter)
    width, height = letter
    margin        = 50
    column_width  = (width - 3 * margin) / 2
    label_height  = 50
    font_size     = 12
    max_per_col   = 10

    c.setFont("Helvetica", font_size)
    left_col  = usernames[:max_per_col]
    right_col = usernames[max_per_col:max_per_col * 2]
    y_left    = height - margin - label_height
    y_right   = height - margin - label_height

    for username in left_col:
        c.setStrokeColor(colors.black)
        c.setFillColor(colors.white)
        c.rect(margin, y_left - label_height, column_width, label_height, fill=1)
        c.setFillColor(colors.black)
        c.drawString(margin + 10, y_left - label_height + 15, username)
        y_left -= label_height + 5

    for username in right_col:
        c.setStrokeColor(colors.black)
        c.setFillColor(colors.white)
        c.rect(margin + column_width + margin, y_right - label_height, column_width, label_height, fill=1)
        c.setFillColor(colors.black)
        c.drawString(margin + column_width + margin + 10, y_right - label_height + 15, username)
        y_right -= label_height + 5

    c.save()

# ── entrypoint ────────────────────────────────────────────────────────────────
# ---------------------------------------------------------
# Practical Evangelism RSVP
# ---------------------------------------------------------

@app.route('/rsvp', methods=['GET', 'POST'])
def rsvp():

    db_path = os.environ.get("DB_PATH", "votestack3.db")

    conn = sqlite3.connect(db_path)

    conn.execute("""
        CREATE TABLE IF NOT EXISTS rsvps (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            first_name TEXT NOT NULL,
            attending TEXT NOT NULL,
            created_at DATETIME DEFAULT CURRENT_TIMESTAMP
        )
    """)

    conn.commit()

    submitted = False

    if request.method == 'POST':

        first_name = request.form.get('first_name', '').strip()
        attending = request.form.get('attending', '').strip()

        if first_name and attending:

            conn.execute(
                "INSERT INTO rsvps (first_name, attending) VALUES (?, ?)",
                (first_name, attending)
            )

            conn.commit()
            submitted = True

    conn.close()

    return render_template(
        'rsvp.html',
        submitted=submitted
    )
# =========================================================
# ADMIN + RSVP SYSTEM
# =========================================================

ADMIN_USERNAME = os.environ.get("ADMIN_USERNAME", "admin")
ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD")

DB_PATH = os.environ.get("DB_PATH", "sfcd.db")


def get_rsvp_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row

    conn.execute("""
        CREATE TABLE IF NOT EXISTS rsvp_events (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            title TEXT NOT NULL,
            event_date TEXT NOT NULL,
            event_time TEXT NOT NULL,
            active INTEGER DEFAULT 1,
            created_at DATETIME DEFAULT CURRENT_TIMESTAMP
        )
    """)

    conn.execute("""
        CREATE TABLE IF NOT EXISTS rsvps (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            event_id INTEGER NOT NULL,
            first_name TEXT NOT NULL,
            attending TEXT NOT NULL,
            created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY(event_id) REFERENCES rsvp_events(id)
        )
    """)

    conn.commit()
    return conn


# ---------------------------------------------------------
# ADMIN LOGIN
# ---------------------------------------------------------

@app.route("/admin_login", methods=["GET", "POST"])
def admin_login():

    error = None

    if request.method == "POST":

        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")

        if not ADMIN_PASSWORD:
            error = "Admin password has not been configured."

        elif (
            hmac.compare_digest(username, ADMIN_USERNAME)
            and hmac.compare_digest(password, ADMIN_PASSWORD)
        ):
            session["admin_logged_in"] = True
            return redirect(url_for("admin_dashboard"))

        else:
            error = "Invalid username or password."

    return render_template("admin_login.html", error=error)


# ---------------------------------------------------------
# ADMIN LOGOUT
# ---------------------------------------------------------

@app.route("/admin_logout")
def admin_logout():
    session.pop("admin_logged_in", None)
    return redirect(url_for("admin_login"))


# ---------------------------------------------------------
# ADMIN DASHBOARD
# ---------------------------------------------------------

@app.route("/admin")
def admin_dashboard():

    if not session.get("admin_logged_in"):
        return redirect(url_for("admin_login"))

    conn = get_rsvp_db()

    events = conn.execute("""
        SELECT
            e.*,

            (
                SELECT COUNT(*)
                FROM rsvps r
                WHERE r.event_id = e.id
                AND r.attending = 'Yes'
            ) AS yes_count,

            (
                SELECT COUNT(*)
                FROM rsvps r
                WHERE r.event_id = e.id
                AND r.attending = 'No'
            ) AS no_count

        FROM rsvp_events e
        ORDER BY e.id DESC
    """).fetchall()

    conn.close()

    return render_template(
        "admin_dashboard.html",
        events=events
    )


# ---------------------------------------------------------
# ADMIN CREATE RSVP EVENT
# ---------------------------------------------------------

@app.route("/admin/rsvp/create", methods=["POST"])
def create_rsvp_event():

    if not session.get("admin_logged_in"):
        return redirect(url_for("admin_login"))

    title = request.form.get("title", "").strip()
    event_date = request.form.get("event_date", "").strip()
    event_time = request.form.get("event_time", "").strip()

    if not title or not event_date or not event_time:
        return redirect(url_for("admin_dashboard"))

    conn = get_rsvp_db()

    cursor = conn.execute("""
        INSERT INTO rsvp_events
        (title, event_date, event_time)
        VALUES (?, ?, ?)
    """, (
        title,
        event_date,
        event_time
    ))

    event_id = cursor.lastrowid

    conn.commit()
    conn.close()

    return redirect(
        url_for(
            "admin_dashboard",
            created=event_id
        )
    )


# ---------------------------------------------------------
# PUBLIC RSVP PAGE
# ---------------------------------------------------------

@app.route("/rsvp/<int:event_id>", methods=["GET", "POST"])
def rsvp_event(event_id):

    conn = get_rsvp_db()

    event = conn.execute(
        "SELECT * FROM rsvp_events WHERE id = ?",
        (event_id,)
    ).fetchone()

    if not event:
        conn.close()
        return "RSVP event not found.", 404

    submitted = False

    if request.method == "POST":

        first_name = request.form.get(
            "first_name", ""
        ).strip()

        attending = request.form.get(
            "attending", ""
        ).strip()

        if first_name and attending in ["Yes", "No"]:

            conn.execute("""
                INSERT INTO rsvps
                (event_id, first_name, attending)
                VALUES (?, ?, ?)
            """, (
                event_id,
                first_name,
                attending
            ))

            conn.commit()

            submitted = True

    conn.close()

    return render_template(
        "rsvp.html",
        event=event,
        submitted=submitted
    )


# ---------------------------------------------------------
# ADMIN RSVP RESULTS
# ---------------------------------------------------------

@app.route("/admin/rsvp/<int:event_id>")
def admin_rsvp_results(event_id):

    if not session.get("admin_logged_in"):
        return redirect(url_for("admin_login"))

    conn = get_rsvp_db()

    event = conn.execute(
        "SELECT * FROM rsvp_events WHERE id = ?",
        (event_id,)
    ).fetchone()

    responses = conn.execute("""
        SELECT *
        FROM rsvps
        WHERE event_id = ?
        ORDER BY created_at DESC
    """, (event_id,)).fetchall()

    conn.close()

    return render_template(
        "admin_rsvp_results.html",
        event=event,
        responses=responses
    )
if __name__ == "__main__":
    app.run(host='0.0.0.0', port=5000, debug=False)
